"""
SignBridge AI - Genuine Excel (.xlsx) Export Service
Generates professional, styled multi-sheet Excel workbooks for user profiles,
confirmed sentence history, and export metadata.
Includes strict spreadsheet formula injection sanitization.
"""

import io
from datetime import datetime, timezone
from typing import List, Dict, Optional, Any
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def sanitize_excel_value(val: Any) -> Any:
    """
    Guards against Spreadsheet Formula Injection (CSV/Formula Injection).
    Treats strings starting with =, +, -, @, \t, or \r as escaped plain text.
    Preserves numbers and booleans.
    """
    if val is None:
        return ""
    if isinstance(val, (int, float, bool)):
        return val
    str_val = str(val)
    if str_val and str_val[0] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + str_val
    return str_val


class ExcelExportService:
    @staticmethod
    def generate_history_workbook(
        user: Optional[Dict[str, Any]],
        sentences: List[Dict[str, Any]],
        storage_source: str = "Local SQLite"
    ) -> io.BytesIO:
        """
        Builds a multi-sheet .xlsx workbook containing:
        1. User Profile Sheet
        2. Sentence History Sheet
        3. Export Information & Disclaimers Sheet
        """
        wb = openpyxl.Workbook()

        # Styles
        primary_header_fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
        secondary_header_fill = PatternFill(start_color="0EA5E9", end_color="0EA5E9", fill_type="solid")
        zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        title_font = Font(name="Segoe UI", size=14, bold=True, color="1E1B4B")
        regular_font = Font(name="Segoe UI", size=10, color="0F172A")
        bold_font = Font(name="Segoe UI", size=10, bold=True, color="0F172A")
        italic_muted_font = Font(name="Segoe UI", size=9, italic=True, color="64748B")

        thin_border_side = Side(border_style="thin", color="CBD5E1")
        cell_border = Border(
            left=thin_border_side,
            right=thin_border_side,
            top=thin_border_side,
            bottom=thin_border_side
        )

        # =========================================================================
        # Sheet 1: User Profile
        # =========================================================================
        ws_profile = wb.active
        ws_profile.title = "User Profile"
        ws_profile.views.sheetView[0].showGridLines = True

        ws_profile.cell(row=1, column=1, value="SignBridge AI — User Profile Details").font = title_font
        ws_profile.cell(row=2, column=1, value=f"Generated on: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M:%S UTC')}").font = italic_muted_font

        profile_headers = ["Field", "Value"]
        for col_idx, h in enumerate(profile_headers, start=1):
            cell = ws_profile.cell(row=4, column=col_idx, value=h)
            cell.font = header_font
            cell.fill = primary_header_fill
            cell.alignment = Alignment(horizontal="left", vertical="center")
            cell.border = cell_border

        profile_rows = [
            ("User Type", "Registered Account" if user else "Guest / Local Session"),
            ("Display Name", sanitize_excel_value(user.get("display_name") if user else "Guest User")),
            ("Email Address", sanitize_excel_value(user.get("email") if user else "guest@local")),
            ("Preferred Language", sanitize_excel_value(user.get("preferred_language", "en") if user else "en")),
            ("Account Created Date", sanitize_excel_value(user.get("created_at") if user else "Current Session")),
            ("Persistence Source", sanitize_excel_value(storage_source))
        ]

        for r_idx, (field, val) in enumerate(profile_rows, start=5):
            c1 = ws_profile.cell(row=r_idx, column=1, value=field)
            c2 = ws_profile.cell(row=r_idx, column=2, value=val)
            c1.font = bold_font
            c2.font = regular_font
            c1.border = cell_border
            c2.border = cell_border
            if r_idx % 2 == 1:
                c1.fill = zebra_fill
                c2.fill = zebra_fill

        # Auto-adjust profile column widths
        for col in ws_profile.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws_profile.column_dimensions[col_letter].width = max(max_len + 5, 20)

        # =========================================================================
        # Sheet 2: Confirmed Sentence History
        # =========================================================================
        ws_sentences = wb.create_sheet(title="Sentence History")
        ws_sentences.views.sheetView[0].showGridLines = True

        ws_sentences.cell(row=1, column=1, value="Confirmed Sign Language Sentence History").font = title_font
        ws_sentences.cell(row=2, column=1, value="Logged human-confirmed sentences formed from recognized Indian Sign Language gestures.").font = italic_muted_font

        sentence_headers = [
            "Record ID",
            "Date & Time (UTC)",
            "Original Recognized Signs",
            "Confirmed Sentence",
            "English Output",
            "Telugu Script Output",
            "Romanized Telugu Output",
            "Selected Language",
            "Recognition Mode"
        ]

        for col_idx, h in enumerate(sentence_headers, start=1):
            cell = ws_sentences.cell(row=4, column=col_idx, value=h)
            cell.font = header_font
            cell.fill = primary_header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = cell_border

        if not sentences:
            empty_cell = ws_sentences.cell(row=5, column=1, value="No confirmed sentences logged yet.")
            empty_cell.font = italic_muted_font
            ws_sentences.merge_cells("A5:I5")
        else:
            for r_idx, item in enumerate(sentences, start=5):
                row_data = [
                    item.get("id", ""),
                    item.get("created_at", ""),
                    item.get("original_signs", ""),
                    item.get("confirmed_sentence", ""),
                    item.get("english_output", "") or "",
                    item.get("telugu_output", "") or "",
                    item.get("romanized_output", "") or "",
                    item.get("selected_language", "en"),
                    item.get("recognition_mode", "custom_isl")
                ]
                for c_idx, val in enumerate(row_data, start=1):
                    cell = ws_sentences.cell(row=r_idx, column=c_idx, value=sanitize_excel_value(val))
                    cell.font = regular_font
                    cell.border = cell_border
                    cell.alignment = Alignment(vertical="center", wrap_text=False)
                    if r_idx % 2 == 1:
                        cell.fill = zebra_fill

        # Auto-adjust sentence column widths
        for col in ws_sentences.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws_sentences.column_dimensions[col_letter].width = min(max(max_len + 4, 15), 50)

        # =========================================================================
        # Sheet 3: Export Information & Audit Metadata
        # =========================================================================
        ws_info = wb.create_sheet(title="Export Information")
        ws_info.views.sheetView[0].showGridLines = True

        ws_info.cell(row=1, column=1, value="SignBridge AI — Export Information & Governance").font = title_font

        info_headers = ["Attribute", "Details"]
        for col_idx, h in enumerate(info_headers, start=1):
            cell = ws_info.cell(row=3, column=col_idx, value=h)
            cell.font = header_font
            cell.fill = secondary_header_fill
            cell.alignment = Alignment(horizontal="left", vertical="center")
            cell.border = cell_border

        info_data = [
            ("Export Generation Timestamp", datetime.now(timezone.utc).isoformat()),
            ("Total Sentence Records Exported", str(len(sentences))),
            ("Application Version", "SignBridge AI v1.0.0"),
            ("Storage Mode", sanitize_excel_value(storage_source)),
            ("Security & Privacy Note", "Passwords, session tokens, secret keys, and biometric data are strictly excluded."),
            ("Formula Injection Protection", "All string fields prefixed with formula symbols are escaped as text literals."),
            ("AI Recognition Transparency", "Sentences are formed from verified gesture vocabulary and grammar rules. Isolated static hand poses are not claimed as full dynamic ISL translation.")
        ]

        for r_idx, (attr, det) in enumerate(info_data, start=4):
            c1 = ws_info.cell(row=r_idx, column=1, value=attr)
            c2 = ws_info.cell(row=r_idx, column=2, value=det)
            c1.font = bold_font
            c2.font = regular_font
            c1.border = cell_border
            c2.border = cell_border
            if r_idx % 2 == 1:
                c1.fill = zebra_fill
                c2.fill = zebra_fill

        for col in ws_info.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws_info.column_dimensions[col_letter].width = min(max(max_len + 5, 20), 80)

        # Output to in-memory bytes
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output


excel_export_service = ExcelExportService()
