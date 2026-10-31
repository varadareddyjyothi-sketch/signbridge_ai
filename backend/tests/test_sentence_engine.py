"""
Backend Tests - Smart Sentence Formation Engine, Grammar, Quick Phrases, and Source Mode
"""

import io
import pytest
import uuid
import openpyxl
from httpx import AsyncClient, ASGITransport

from backend.app.main import app
from backend.app.services.sentence_service import (
    SignToMeaningMapper,
    GrammarEngine,
    sentence_translation_service,
)
from backend.app.services.phrase_library_service import phrase_library_service


# =========================================================================
# 1. Sign-to-Meaning Mapping Tests
# =========================================================================

def test_sign_to_meaning_supported_and_unsupported():
    """Verify supported signs map to approved meanings and unsupported signs are marked correctly."""
    mapper = SignToMeaningMapper()

    # Supported sign
    mapping_hello = mapper.map_token("HELLO")
    assert mapping_hello["supported"] is True
    assert mapping_hello["english"].lower() == "hello"
    assert "నమస్కారం" in mapping_hello["telugu"]

    # Supported two-hand sign
    mapping_namaste = mapper.map_token("ISL_NAMASTE")
    assert mapping_namaste["supported"] is True
    assert "namaste" in mapping_namaste["english"].lower()

    # Unsupported sign - transparent reporting, no fake fabrication
    mapping_fake = mapper.map_token("UNSUPPORTED_RANDOM_GESTURE_99")
    assert mapping_fake["supported"] is False
    assert "Unrecognized" in mapping_fake["english"] or "Unsupported" in mapping_fake["english"]
    assert "గుర్తించబడని" in mapping_fake["telugu"]


# =========================================================================
# 2. Grammar Engine & Sentence Intent Tests
# =========================================================================

def test_grammar_questions_punctuation():
    """Questions should be classified with question intent and receive a question mark (?)."""
    mapper = SignToMeaningMapper()
    grammar = GrammarEngine(mapper)

    # WHERE BUS STOP
    res_where = grammar.analyze_sequence(["WHERE", "BUS", "STOP"])
    assert res_where["intent"] == "question"
    assert res_where["punctuation"] == "?"
    assert res_where["english"].endswith("?")
    assert "where" in res_where["english"].lower()

    # WHO YOU
    res_who = grammar.analyze_sequence(["WHO", "YOU"])
    assert res_who["intent"] == "question"
    assert res_who["english"].endswith("?")

    # HOW YOU
    res_how = grammar.analyze_sequence(["HOW", "YOU"])
    assert res_how["intent"] == "question"
    assert res_how["english"].endswith("?")


def test_grammar_requests_and_greetings():
    """Requests receive polite period/statement ending, greetings receive exclamation mark."""
    mapper = SignToMeaningMapper()
    grammar = GrammarEngine(mapper)

    # PLEASE HELP
    res_req = grammar.analyze_sequence(["PLEASE", "HELP"])
    assert res_req["intent"] == "request"
    assert "help" in res_req["english"].lower()
    assert res_req["english"].endswith(".")

    # HELLO
    res_greet = grammar.analyze_sequence(["HELLO"])
    assert res_greet["intent"] == "greeting"
    assert res_greet["english"].endswith("!")


def test_grammar_numeral_noun_inflection():
    """Numeral + Noun combinations properly inflect singular vs plural."""
    mapper = SignToMeaningMapper()
    grammar = GrammarEngine(mapper)

    # 1 book (singular)
    res_1_book = grammar.analyze_sequence(["ISL_1", "ISL_BOOK"])
    assert "one study book." in res_1_book["english"].lower() or "one book." in res_1_book["english"].lower()

    # 2 books (plural)
    res_2_book = grammar.analyze_sequence(["ISL_2", "ISL_BOOK"])
    assert "two study books." in res_2_book["english"].lower() or "two books." in res_2_book["english"].lower()


def test_grammar_negation_and_affirmation():
    """Polarity signs (YES / NO) produce accurate intent classifications."""
    mapper = SignToMeaningMapper()
    grammar = GrammarEngine(mapper)

    res_yes = grammar.analyze_sequence(["YES"])
    assert res_yes["intent"] == "affirmation"
    assert "Yes" in res_yes["english"]

    res_no = grammar.analyze_sequence(["NO"])
    assert res_no["intent"] == "negation"
    assert "No" in res_no["english"]


def test_sentence_service_breakdown_and_disambiguation():
    """Sentence translation service returns token breakdowns and disambiguation choices."""
    result = sentence_translation_service.translate_signs_sequence(["PLEASE", "WATER"])
    assert result["sentence_intent"] == "request"
    assert len(result["meaning_breakdown"]) == 2
    assert result["meaning_breakdown"][0]["token"] == "PLEASE"
    assert result["meaning_breakdown"][1]["token"] == "WATER"

    # Disambiguation for brief ambiguous gesture
    disambig_res = sentence_translation_service.translate_signs_sequence(["ISL_5"])
    assert len(disambig_res["disambiguation_options"]) >= 2


# =========================================================================
# 3. Quick Phrase Library Service Tests (Mode 2)
# =========================================================================

def test_quick_phrase_categories_and_counts():
    """Verify that all 8 categories (A-H) exist and total 63 reviewed phrases."""
    categories = phrase_library_service.get_categories()
    assert len(categories) == 8
    cat_ids = [c["id"] for c in categories]
    assert "greetings" in cat_ids
    assert "questions" in cat_ids
    assert "conversation" in cat_ids
    assert "travel" in cat_ids
    assert "needs" in cat_ids
    assert "education" in cat_ids
    assert "health_emergency" in cat_ids
    assert "responses" in cat_ids

    all_phrases = phrase_library_service.get_phrases()
    assert len(all_phrases) == 63


def test_quick_phrase_search_and_emergency():
    """Search by English and Telugu, and check emergency flags."""
    # Search water
    water_matches = phrase_library_service.get_phrases(search="water")
    assert any("water" in p["english"].lower() for p in water_matches)

    # Search Telugu
    telugu_matches = phrase_library_service.get_phrases(search="ఎలా ఉన్నారు")
    assert len(telugu_matches) > 0
    assert "How are you?" in telugu_matches[0]["english"]

    # Category filter: health_emergency
    emergency_phrases = phrase_library_service.get_phrases(category="health_emergency")
    assert len(emergency_phrases) >= 5
    assert all(p["is_emergency"] is True for p in emergency_phrases)
    ambulance_phrase = next((p for p in emergency_phrases if "ambulance" in p["english"].lower()), None)
    assert ambulance_phrase is not None


# =========================================================================
# 4. API Endpoints Tests (Mode 1 & Mode 2)
# =========================================================================

@pytest.mark.anyio
async def test_api_translate_sentence():
    """Test POST /translate/sentence live formation."""
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as ac:
        payload = {"tokens": ["HELLO", "HOW", "YOU"]}
        res = await ac.post("/translate/sentence", json=payload)
        assert res.status_code == 200
        data = res.json()
        assert data["sentence_intent"] in ["greeting", "question"]
        assert "telugu" in data
        assert "telugu_romanized" in data
        assert len(data["meaning_breakdown"]) == 3


@pytest.mark.anyio
async def test_api_quick_phrase_endpoints():
    """Test GET /translate/categories and GET /translate/phrases."""
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as ac:
        # Categories
        res_cats = await ac.get("/translate/categories")
        assert res_cats.status_code == 200
        cats = res_cats.json()
        assert len(cats) == 8

        # Filtered phrases
        res_phrases = await ac.get("/translate/phrases?category=basic_needs")
        assert res_phrases.status_code == 200
        data = res_phrases.json()
        assert data["total"] > 0
        assert all(p["category"] in ("needs", "basic_needs") for p in data["phrases"])


# =========================================================================
# 5. History & Excel Export with Source Mode
# =========================================================================

@pytest.mark.anyio
async def test_sentence_history_and_excel_export_with_source_mode():
    """Verify storing sentence with source_mode and exporting to Excel with Source Mode column."""
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as ac:
        email = f"mode_test_{uuid.uuid4().hex[:8]}@example.com"
        res_reg = await ac.post("/auth/register", json={
            "email": email,
            "password": "Password123!",
            "display_name": "Mode Tester",
            "preferred_language": "en"
        })
        token = res_reg.cookies.get("signbridge_session")
        headers = {"Authorization": f"Bearer {token}"}

        # 1. Save sentence from Mode 1 (Camera Recognition)
        res_cam = await ac.post("/history/sentence", headers=headers, json={
            "original_signs": ["HELLO", "THANK_YOU"],
            "confirmed_sentence": "Hello, thank you!",
            "telugu_translation": "నమస్కారం, ధన్యవాదాలు!",
            "romanized_telugu": "Namaskaram, dhanyavaadaalu!",
            "source_mode": "camera_recognition"
        })
        assert res_cam.status_code == 200
        assert res_cam.json()["source_mode"] == "camera_recognition"

        # 2. Save sentence from Mode 2 (Quick Phrase Selection)
        res_quick = await ac.post("/history/sentence", headers=headers, json={
            "original_signs": ["QUICK_PHRASE:help_me"],
            "confirmed_sentence": "Please help me.",
            "telugu_translation": "దయచేసి నాకు సహాయం చేయండి.",
            "romanized_telugu": "Dayachesi naaku sahaayam cheyandi.",
            "source_mode": "quick_phrase"
        })
        assert res_quick.status_code == 200
        assert res_quick.json()["source_mode"] == "quick_phrase"

        # 3. Retrieve history and verify source_mode
        res_hist = await ac.get("/history/sentences", headers=headers)
        assert res_hist.status_code == 200
        hist_data = res_hist.json()
        items = hist_data.get("items", [])
        assert len(items) >= 2
        sources = [item["source_mode"] for item in items]
        assert "camera_recognition" in sources
        assert "quick_phrase" in sources

        # 4. Export to Excel and verify 'Source Mode' column
        res_export = await ac.get("/history/export/excel", headers=headers)
        assert res_export.status_code == 200
        wb = openpyxl.load_workbook(filename=io.BytesIO(res_export.content))
        sentence_sheet = wb["Sentence History"]
        headers_row = [cell.value for cell in sentence_sheet[4]]
        assert "Source Mode" in headers_row

        # Verify values in Source Mode column
        mode_idx = headers_row.index("Source Mode")
        row_modes = [sentence_sheet.cell(row=r, column=mode_idx + 1).value for r in range(5, sentence_sheet.max_row + 1)]
        assert "Camera Recognition" in row_modes
        assert "Quick Phrase Selection" in row_modes


@pytest.mark.anyio
async def test_duplicate_sentence_suppression_and_single_deletion():
    """Verify rapid duplicate submissions are suppressed and individual deletion works."""
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as ac:
        email = f"dup_tester_{uuid.uuid4().hex[:8]}@example.com"
        res_reg = await ac.post("/auth/register", json={
            "email": email,
            "password": "Password123!",
            "display_name": "Deduplication Tester",
            "preferred_language": "en"
        })
        token = res_reg.cookies.get("signbridge_session")
        headers = {"Authorization": f"Bearer {token}"}

        # 1. Post identical sentence twice rapidly
        payload = {
            "original_signs": ["PLEASE", "HELP"],
            "confirmed_sentence": "Please help me.",
            "telugu_translation": "దయచేసి నాకు సహాయం చేయండి.",
            "source_mode": "camera_recognition"
        }
        res_1 = await ac.post("/history/sentence", headers=headers, json=payload)
        res_2 = await ac.post("/history/sentence", headers=headers, json=payload)

        assert res_1.status_code == 200
        assert res_2.status_code == 200
        # Idempotency: the second submission returns the existing record ID
        assert res_1.json()["id"] == res_2.json()["id"]

        # 2. Check history count is 1 (not duplicated)
        res_hist = await ac.get("/history/sentences", headers=headers)
        items = res_hist.json().get("items", [])
        assert len([i for i in items if i["confirmed_sentence"] == "Please help me."]) == 1

        # 3. Individual sentence deletion
        record_id = res_1.json()["id"]
        res_del = await ac.delete(f"/history/sentence/{record_id}", headers=headers)
        assert res_del.status_code == 200
        assert res_del.json()["deleted_id"] == record_id

        # 4. Verify record is now gone
        res_hist_after = await ac.get("/history/sentences", headers=headers)
        items_after = res_hist_after.json().get("items", [])
        assert not any(i["id"] == record_id for i in items_after)
