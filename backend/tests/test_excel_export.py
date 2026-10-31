"""Backend API tests for Excel export and formula injection prevention."""

import io
import uuid

import mongomock
import openpyxl
import pytest
from httpx import ASGITransport, AsyncClient

from backend.app.database.mongo_client import set_mock_mongo_db
from backend.app.main import app
from backend.app.services.excel_export_service import sanitize_excel_value


@pytest.fixture(autouse=True)
def isolated_mongo():
    client = mongomock.MongoClient()
    set_mock_mongo_db(client["signbridge_ai"])
    yield
    set_mock_mongo_db(None)
    client.close()


def test_formula_injection_sanitization():
    assert sanitize_excel_value("=1+1") == "'=1+1"
    assert sanitize_excel_value("+cmd|'/c calc'!A1") == "'+cmd|'/c calc'!A1"
    assert sanitize_excel_value("-5*10") == "'-5*10"
    assert sanitize_excel_value("@SUM(A1:A10)") == "'@SUM(A1:A10)"
    assert sanitize_excel_value("\tTAB_INJECT") == "'\tTAB_INJECT"
    assert sanitize_excel_value("Safe string") == "Safe string"
    assert sanitize_excel_value(123) == 123
    assert sanitize_excel_value(None) == ""


@pytest.mark.anyio
async def test_excel_export_endpoints():
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        email = f"excel_user_{uuid.uuid4().hex}@example.com"
        registration = await client.post(
            "/auth/register",
            json={
                "email": email,
                "password": "Password123!",
                "display_name": "Excel Tester",
                "preferred_language": "te",
            },
        )
        assert registration.status_code == 200

        export_empty = await client.get("/history/export/excel")
        assert export_empty.status_code == 200
        assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in (
            export_empty.headers.get("content-type", "")
        )
        workbook_empty = openpyxl.load_workbook(io.BytesIO(export_empty.content))
        assert {"User Profile", "Sentence History", "Export Information"}.issubset(
            workbook_empty.sheetnames
        )
        info_rows = list(workbook_empty["Export Information"].iter_rows(values_only=True))
        assert any("Total Sentence Records" in str(row[0]) for row in info_rows)

        malicious_input = "=cmd|' /C calc'!A0"
        save_response = await client.post(
            "/history/sentence",
            json={
                "original_signs": ["YES", "HELLO"],
                "confirmed_sentence": malicious_input,
                "english_output": "Formula Test",
                "telugu_output": "సూత్ర పరీక్ష",
                "romanized_telugu": "Soothra Pareeksha",
                "selected_language": "te",
                "recognition_mode": "sentence_builder",
            },
        )
        assert save_response.status_code == 200

        export_response = await client.get("/history/export/excel")
        assert export_response.status_code == 200
        workbook = openpyxl.load_workbook(io.BytesIO(export_response.content))
        rows = list(workbook["Sentence History"].iter_rows(values_only=True))
        header_index = next(i for i, row in enumerate(rows) if row and "Record ID" in row)
        headers = rows[header_index]
        assert "Confirmed Sentence" in headers
        values = [
            row[headers.index("Confirmed Sentence")]
            for row in rows[header_index + 1:]
            if row and row[headers.index("Confirmed Sentence")] is not None
        ]
        assert f"'{malicious_input}" in values
